"""建联台账（Excel）自动生成。

从 SQLite 里把「建联行动」相关的账号导出成一个可直接复制粘贴的 xlsx：
  - Sheet「建联台账」：明细（品牌 / 账号 / 粉丝 / 发送时间 / 状态 / 回复内容 / 回复方向 / 备注）
  - Sheet「汇总」：总发送、总回复、回复率、各回复方向占比
  - Sheet「已发送去重名单」：所有已发出账号，下次发送前人工核对、绝不重复

每次发送 / 检测回复后调用 export_tracker() 刷新一次，Excel 永远是最新进度。
"""
import time
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------- 回复方向自动分类（中英双语关键词） ----------
# 判断顺序很重要：先判负向（"not interested" 里含 "interest"），再正向，再中性。
NEGATIVE_KW = [
    "not interest", "not interested", "no interest", "no thanks", "no need",
    "dont need", "don't need", "do not need", "not looking", "not at this time",
    "not a fit", "no budget", "already work", "already have", "not for us",
    "stop", "remove", "unsubscribe", "not currently",
    "不需要", "不感兴趣", "谢谢不用", "暂时不做", "目前没有", "没有预算",
    "拒绝", "不考虑", "不做了", "暂时不需要",
]
POSITIVE_KW = [
    "interested", "interest", "sure", "yes", "great", "love to", "send",
    "media kit", "deck", "portfolio", "budget", "price", "quote", "pricing",
    "wechat", "whatsapp", "email", "call", "schedule", "meeting", "connect",
    "partner", "work together", "let's talk", "lets talk", "tell me more",
    "感兴趣", "合作", "预算", "报价", "微信", "邮箱", "聊聊", "聊一下",
    "发资料", "详谈", "可以", "怎么合作", "了解一下你们",
]
NEUTRAL_KW = [
    "thanks", "thank you", "consider", "maybe", "later", "get back",
    "will review", "not now", "looking into", "keep in touch",
    "谢谢", "考虑", "以后", "稍后", "看看", "先了解",
]

DIRECTION_LABEL = {
    "hot": "正向·感兴趣",
    "warm": "中性·在考虑",
    "cold": "负向·拒绝",
    "pending": "待判断",
}
STATUS_LABEL = {
    "drafted": "待发送",
    "sent": "已发送·待回复",
    "replied": "已回复",
    "second_touched": "已回复·已跟进",
    "followed_up": "已跟进·待回复",
    "failed": "发送失败",
    "rejected": "已淘汰(错号)",
    "new": "待补全",
    "enriched": "已补全",
    "qualified": "待写文案",
    "blocked": "风控暂停",
}

# 台账明细只关心「建联行动」这些状态
ACTION_STATUSES = ("drafted", "sent", "replied", "followed_up",
                   "second_touched", "failed")

SENT_STATUSES = ("sent", "replied", "followed_up", "second_touched")


def classify_direction(reply_text: Optional[str]) -> str:
    """根据回复内容自动归类：hot / warm / cold / pending。"""
    if not reply_text:
        return "pending"
    t = reply_text.lower()
    for kw in NEGATIVE_KW:
        if kw in t:
            return "cold"
    for kw in POSITIVE_KW:
        if kw in t:
            return "hot"
    for kw in NEUTRAL_KW:
        if kw in t:
            return "warm"
    return "pending"


def _fmt_ts(ts) -> str:
    if not ts:
        return ""
    return time.strftime("%Y-%m-%d %H:%M", time.localtime(ts))


def _direction_of(lead: Dict) -> str:
    """回复方向：优先人工 intent（hot/warm/cold），否则用自动分类结果。"""
    intent = (lead.get("intent") or "").strip()
    if intent in ("hot", "warm", "cold"):
        return intent
    direction = (lead.get("direction") or "").strip()
    if direction in ("hot", "warm", "cold", "pending"):
        return direction
    # 兜底：实时再算一次
    return classify_direction(lead.get("reply_snippet"))


def _is_replied(lead: Dict) -> bool:
    return lead.get("status") in ("replied", "second_touched")


def export_tracker(store, path: Optional[Path] = None) -> Path:
    """生成 / 刷新建联台账 xlsx，返回文件路径。"""
    if path is None:
        from .config import DATA_DIR
        path = Path(DATA_DIR) / "建联台账.xlsx"

    leads = store.all()
    action = [l for l in leads if l.get("status") in ACTION_STATUSES]
    # 排序：已发送的按发送时间倒序（最新在前），待发的按粉丝数倒序
    action.sort(key=lambda l: (
        0 if l.get("status") in SENT_STATUSES + ("failed",) else 1,
        -(l.get("sent_at") or 0) if l.get("status") in SENT_STATUSES + ("failed",)
        else -(l.get("followers") or 0),
    ))

    sent_leads = [l for l in leads if l.get("status") in SENT_STATUSES]
    replied_leads = [l for l in sent_leads if _is_replied(l)]

    wb = Workbook()

    # ---------- Sheet 1：建联台账 ----------
    ws = wb.active
    ws.title = "建联台账"
    headers = ["序号", "品牌", "IG账号", "粉丝数", "状态",
               "发送时间", "回复时间", "回复内容", "回复方向", "邮箱", "社交线索", "备注"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    dir_fill = {
        "hot": PatternFill("solid", fgColor="C6EFCE"),
        "warm": PatternFill("solid", fgColor="FFEB9C"),
        "cold": PatternFill("solid", fgColor="FFC7CE"),
        "pending": PatternFill("solid", fgColor="EDEDED"),
    }
    dir_font = {
        "hot": Font(color="006100"),
        "warm": Font(color="9C6500"),
        "cold": Font(color="9C0006"),
        "pending": Font(color="595959"),
    }

    for i, l in enumerate(action, 1):
        d = _direction_of(l)
        ws.append([
            i,
            l.get("full_name") or "",
            "@" + (l.get("username") or ""),
            l.get("followers") or 0,
            STATUS_LABEL.get(l.get("status"), l.get("status")),
            _fmt_ts(l.get("sent_at")),
            _fmt_ts(l.get("replied_at")),
            (l.get("reply_snippet") or "")[:500],
            DIRECTION_LABEL.get(d, d),
            (l.get("email") or "")[:60],
            (l.get("social_links") or "")[:200],
            (l.get("notes") or "")[:300],
        ])
        row = ws.max_row
        # 回复方向着色
        ws.cell(row=row, column=9).fill = dir_fill.get(d, dir_fill["pending"])
        ws.cell(row=row, column=9).font = dir_font.get(d, dir_font["pending"])
        # 已回复的行高亮品牌名
        if _is_replied(l):
            ws.cell(row=row, column=2).font = Font(bold=True, color="1F4E78")

    widths = [6, 20, 24, 10, 16, 16, 16, 40, 14, 26, 34, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # ---------- Sheet 2：汇总 ----------
    ws2 = wb.create_sheet("汇总")
    total_sent = len(sent_leads)
    total_replied = len(replied_leads)
    rate = round(total_replied / total_sent * 100, 1) if total_sent else 0.0

    # 各方向统计（只统计已回复的）
    dir_count = {"hot": 0, "warm": 0, "cold": 0, "pending": 0}
    for l in replied_leads:
        d = _direction_of(l)
        if d in dir_count:
            dir_count[d] += 1

    ws2.append(["KPI", "数值"])
    ws2.append(["总发送数", total_sent])
    ws2.append(["总回复数", total_replied])
    ws2.append(["回复率", f"{rate}%"])
    ws2.append(["待发送（队列中）", sum(1 for l in action if l.get("status") == "drafted")])
    ws2.append(["发送失败", sum(1 for l in action if l.get("status") == "failed")])
    ws2.append([])
    ws2.append(["回复方向分布", "数量", "占已回复比"])
    for d, label in DIRECTION_LABEL.items():
        c = dir_count[d]
        pct = round(c / total_replied * 100, 1) if total_replied else 0.0
        ws2.append([label, c, f"{pct}%"])

    ws2["A1"].font = Font(bold=True, color="FFFFFF")
    ws2["A1"].fill = header_fill
    ws2["B1"].font = Font(bold=True, color="FFFFFF")
    ws2["B1"].fill = header_fill
    for r in range(2, 8):
        ws2.cell(row=r, column=1).font = Font(bold=True)
    ws2.cell(row=8, column=1).font = Font(bold=True)
    ws2.cell(row=8, column=2).font = Font(bold=True)
    ws2.cell(row=8, column=3).font = Font(bold=True)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 14

    # ---------- Sheet 3：已发送去重名单 ----------
    ws3 = wb.create_sheet("已发送去重名单")
    ws3.append(["序号", "品牌", "IG账号", "发送时间", "状态"])
    for c in range(1, 6):
        cell = ws3.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sent_sorted = sorted(sent_leads, key=lambda l: -(l.get("sent_at") or 0))
    for i, l in enumerate(sent_sorted, 1):
        ws3.append([
            i,
            l.get("full_name") or "",
            "@" + (l.get("username") or ""),
            _fmt_ts(l.get("sent_at")),
            STATUS_LABEL.get(l.get("status"), l.get("status")),
        ])
    ws3.append([])
    ws3.append([f"⚠ 以上 {len(sent_sorted)} 个账号已发送过，下次发送前务必核对，绝不重复。"])
    ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, color="9C0006")
    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 24
    ws3.column_dimensions["C"].width = 26
    ws3.column_dimensions["D"].width = 18
    ws3.column_dimensions["E"].width = 16
    ws3.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
